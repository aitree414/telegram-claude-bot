"""
08 Collective — Working Schedule Generator
============================================
Generates HTML + XLSX production schedules for exhibitions/events.

Phases: Load-in → Assembly → Tech Rehearsal → Show → Strike
"""

from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def generate_schedule_html(project, phases, output_path):
    """Generate an HTML working schedule.

    phases: list of {
        name, date_start, date_end, tasks: [{time, activity, crew, note}]
    }
    """
    lines = []
    lines.append('<!DOCTYPE html><html lang="zh-TW"><head><meta charset="UTF-8">')
    lines.append('<title>Working Schedule — ' + project + '</title>')
    lines.append('<style>')
    lines.append('  body { font-family: "Helvetica Neue", "Microsoft JhengHei", sans-serif; background: #0a1628; color: #caf0f8; padding: 20px; }')
    lines.append('  h1 { color: #00b4d8; border-bottom: 2px solid #00b4d8; padding-bottom: 8px; }')
    lines.append('  .phase { background: #0d1f3c; border: 1px solid #00b4d8; border-radius: 4px; margin: 16px 0; padding: 12px; }')
    lines.append('  .phase h2 { color: #f77f00; margin: 0 0 8px 0; font-size: 14px; }')
    lines.append('  .phase .date { color: #90e0ef; font-size: 11px; }')
    lines.append('  table { width: 100%; border-collapse: collapse; font-size: 11px; }')
    lines.append('  th { background: #1a3a5c; color: #00b4d8; padding: 6px; text-align: left; border: 1px solid #1a3a5c; }')
    lines.append('  td { padding: 4px 6px; border: 1px solid #1a3a5c; color: #caf0f8; }')
    lines.append('  tr:nth-child(even) { background: #0a1628; }')
    lines.append('  .note-warn { color: #f77f00; }')
    lines.append('</style></head><body>')
    lines.append(f'<h1>{project} — Working Schedule 工作進度表</h1>')

    for phase in phases:
        lines.append(f'<div class="phase">')
        lines.append(f'<h2>{phase["name"]}</h2>')
        lines.append(f'<div class="date">{phase["date_start"]} → {phase["date_end"]}</div>')
        if phase.get("tasks"):
            lines.append('<table><tr><th>Time</th><th>Activity</th><th>Crew</th><th>Note</th></tr>')
            for t in phase["tasks"]:
                lines.append(f'<tr><td>{t.get("time","")}</td><td>{t.get("activity","")}</td>'
                           f'<td>{t.get("crew","")}</td><td>{t.get("note","")}</td></tr>')
            lines.append('</table>')
        lines.append('</div>')

    lines.append('</body></html>')

    html = '\n'.join(lines)
    with open(output_path, 'w') as f:
        f.write(html)
    return output_path


def generate_schedule_xlsx(project, phases, output_path):
    """Generate an XLSX working schedule."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    ws.cell(row=1, column=1, value=f"{project} — Working Schedule").font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 30

    row = 3
    for phase in phases:
        ws.cell(row=row, column=1, value=phase["name"]).font = Font(bold=True, color='1a1a2e')
        ws.cell(row=row, column=1).fill = PatternFill('solid', fgColor='e8edf2')
        ws.cell(row=row, column=2, value=f"{phase['date_start']} → {phase['date_end']}")
        ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor='e8edf2')
        row += 1

        for t in phase.get("tasks", []):
            ws.cell(row=row, column=1, value=t.get("time",""))
            ws.cell(row=row, column=2, value=t.get("duration",""))
            ws.cell(row=row, column=3, value=t.get("activity",""))
            ws.cell(row=row, column=4, value=t.get("crew",""))
            ws.cell(row=row, column=5, value=t.get("note",""))
            row += 1
        row += 1

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    phases = [
        {
            "name": "Load-in / 進場",
            "date_start": "2026-07-01 09:00",
            "date_end": "2026-07-01 18:00",
            "tasks": [
                {"time": "09:00-10:00", "activity": "Material delivery", "crew": "4", "note": "Truck #1"},
                {"time": "10:00-12:00", "activity": "Truss assembly", "crew": "4", "note": ""},
                {"time": "13:00-18:00", "activity": "Structure setup", "crew": "6", "note": ""},
            ]
        },
        {
            "name": "Tech Rehearsal / 技術排练",
            "date_start": "2026-07-02 10:00",
            "date_end": "2026-07-02 18:00",
            "tasks": [
                {"time": "10:00-12:00", "activity": "Lighting focus", "crew": "2", "note": ""},
                {"time": "13:00-15:00", "activity": "Sound check", "crew": "2", "note": ""},
                {"time": "15:00-18:00", "activity": "Full rehearsal", "crew": "8", "note": ""},
            ]
        },
        {
            "name": "Show Day / 活動日",
            "date_start": "2026-07-03 08:00",
            "date_end": "2026-07-03 22:00",
            "tasks": [
                {"time": "08:00-10:00", "activity": "Final prep", "crew": "4", "note": ""},
                {"time": "10:00-18:00", "activity": "Show open", "crew": "4", "note": ""},
                {"time": "18:00-22:00", "activity": "Strike", "crew": "6", "note": ""},
            ]
        },
    ]

    html_out = "/Users/aitree414/telegram-claude-bot/08_AI_Tools/demo_schedule.html"
    generate_schedule_html("Demo Exhibition", phases, html_out)
    print(f"✓ Schedule HTML: {html_out}")
